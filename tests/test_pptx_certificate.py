"""
tests/test_pptx_certificate.py
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pytest
from pptx import Presentation

import pptx_certificate as pc
import qr_utils


def test_fill_certificate_text_replaces_all_three_fields():
    prs = Presentation(pc.DEFAULT_TEMPLATE_PATH)
    pc.fill_certificate_text(
        prs, name="Test Student", certificate_number="DVA-APIDS-2026-000999", completion_date="01-01-2026"
    )
    slide = prs.slides[0]
    all_text = "\n".join(
        "".join(r.text for r in p.runs)
        for shape in slide.shapes if shape.has_text_frame
        for p in shape.text_frame.paragraphs
    )
    assert "Test Student" in all_text
    assert "DVA-APIDS-2026-000999" in all_text
    assert "01-01-2026" in all_text
    # placeholder brackets should be gone, not just appended alongside
    assert "[                    ]" not in all_text


def test_fill_certificate_text_raises_on_missing_template_field(tmp_path):
    # A blank presentation has none of the expected labels — must fail loudly,
    # not silently produce a certificate with empty fields.
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    with pytest.raises(pc.TemplateFieldNotFound):
        pc.fill_certificate_text(prs, name="X", certificate_number="Y", completion_date="Z")


def test_render_certificate_pdf_end_to_end(tmp_path):
    qr_path = tmp_path / "qr.png"
    qr_path.write_bytes(qr_utils.make_qr_image_bytes("https://example.com/verify/DVA-APIDS-2026-000999"))
    out_path = tmp_path / "cert.pdf"

    pc.render_certificate_pdf(
        name="Test Student",
        certificate_number="DVA-APIDS-2026-000999",
        completion_date="01-01-2026",
        qr_png_path=str(qr_path),
        output_pdf_path=str(out_path),
    )
    assert out_path.exists()
    assert out_path.stat().st_size > 10_000  # sanity: not an empty/broken PDF


def test_qr_encodes_the_correct_verify_url():
    url = qr_utils.verification_url("DVA-APIDS-2026-000123", base_url="https://certs.example.com")
    assert url == "https://certs.example.com/verify/DVA-APIDS-2026-000123"


@pytest.mark.parametrize("course", ["APIDS", "APDA"])
def test_both_course_templates_are_real_distinct_files_and_fill_correctly(course):
    """Regression test: at one point 'APIDA'/'APIDIA'/'APIDS' all pointed at
    byte-identical files (a copy-paste bug), so the course dropdown looked
    like it worked but always rendered the same design. Each course option
    must map to its own real template file."""
    template_path = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "assets", "templates", f"certificate_{course}_blank.pptx",
    )
    assert os.path.exists(template_path), f"Missing template for course {course}"

    prs = Presentation(template_path)
    pc.fill_certificate_text(
        prs, name=f"Test Student {course}", certificate_number=f"DVA-{course}-2026-000001", completion_date="01-01-2026"
    )
    slide = prs.slides[0]
    all_text = "\n".join(
        "".join(r.text for r in p.runs)
        for shape in slide.shapes if shape.has_text_frame
        for p in shape.text_frame.paragraphs
    )
    assert f"Test Student {course}" in all_text
    assert f"DVA-{course}-2026-000001" in all_text


def test_apids_and_apda_templates_are_not_byte_identical():
    base = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "templates")
    with open(os.path.join(base, "certificate_APIDS_blank.pptx"), "rb") as f:
        apids_bytes = f.read()
    with open(os.path.join(base, "certificate_APDA_blank.pptx"), "rb") as f:
        apda_bytes = f.read()
    assert apids_bytes != apda_bytes
